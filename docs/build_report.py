from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
IMAGES = DOCS / "screenshots"
OUTPUT = DOCS / "Bao-cao-do-an-SkinFood.docx"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"


def font(size=28, bold=False):
    filename = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / filename), size)


def wrap(draw, text, fnt, width):
    words, lines, current = text.split(), [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def box(draw, xy, title, rows, fill="#FFFFFF"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline="#4F6B85", width=3)
    draw.rectangle((x1, y1, x2, y1 + 48), fill="#17365D")
    draw.text((x1 + 14, y1 + 10), title, font=font(23, True), fill="white")
    y = y1 + 60
    for row in rows:
        draw.text((x1 + 14, y), row, font=font(19), fill="#202830")
        y += 29


def arrow(draw, start, end, label=""):
    draw.line((*start, *end), fill="#64788C", width=4)
    x, y = end
    draw.polygon([(x, y), (x - 15, y - 8), (x - 15, y + 8)], fill="#64788C")
    if label:
        mx, my = (start[0] + end[0]) // 2, (start[1] + end[1]) // 2
        draw.text((mx - 20, my - 24), label, font=font(17, True), fill="#17365D")


def build_erd():
    image = Image.new("RGB", (1800, 1150), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 30), "Mô hình quan hệ dữ liệu chính", font=font(36, True), fill="#111111")
    entities = {
        "DanhMuc": ((60, 120, 390, 300), ["PK MaDM", "TenDM"]),
        "LoaiSP": ((520, 120, 850, 330), ["PK MaLoai", "FK MaDM", "TenLoai"]),
        "SanPham": ((990, 90, 1410, 390), ["PK MaSP", "FK MaDM", "FK MaLoai", "FK MaTH", "TenSP", "GiaBan", "SoLuongTon"]),
        "ThuongHieu": ((1460, 120, 1740, 300), ["PK MaTH", "TenTH", "QuocGia"]),
        "NguoiDung": ((60, 540, 390, 780), ["PK MaND", "HoTen", "SoDienThoai", "DiaChi"]),
        "TaiKhoan": ((60, 850, 390, 1070), ["PK MaTK", "FK MaND", "TenDangNhap", "VaiTro", "TrangThai"]),
        "DonHang": ((550, 500, 960, 820), ["PK MaDH", "FK MaND", "FK MaKM", "TongTien", "PhuongThuc", "TrangThaiTT", "MaThamChieu"]),
        "ChiTietDonHang": ((1080, 540, 1480, 790), ["PK FK MaDH", "PK FK MaSP", "SoLuong", "DonGia"]),
        "KhuyenMai": ((1450, 850, 1740, 1080), ["PK MaKM", "MaCode", "LoaiGiamGia", "GiaTriGiam"]),
    }
    for title, (xy, rows) in entities.items():
        box(draw, xy, title, rows, "#F7FAFC")
    arrow(draw, (390, 210), (520, 210), "1 n")
    arrow(draw, (850, 220), (990, 220), "1 n")
    arrow(draw, (1460, 220), (1410, 220), "1 n")
    arrow(draw, (390, 650), (550, 650), "1 n")
    arrow(draw, (960, 650), (1080, 650), "1 n")
    arrow(draw, (1410, 390), (1280, 540), "1 n")
    arrow(draw, (220, 780), (220, 850), "1 1")
    arrow(draw, (1450, 960), (900, 820), "1 n")
    path = IMAGES / "04-erd.png"
    image.save(path)
    return path


def build_use_case():
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 30), "UML Use Case của hệ thống", font=font(36, True), fill="#111111")
    draw.rounded_rectangle((330, 100, 1710, 990), radius=20, outline="#8294A6", width=4)
    draw.text((355, 120), "Website SkinFood", font=font(25, True), fill="#17365D")
    actors = [(135, 320, "Khách hàng"), (135, 760, "Quản trị viên")]
    for x, y, label in actors:
        draw.ellipse((x - 25, y - 95, x + 25, y - 45), outline="#1C2833", width=4)
        draw.line((x, y - 45, x, y + 30), fill="#1C2833", width=4)
        draw.line((x - 40, y - 10, x + 40, y - 10), fill="#1C2833", width=4)
        draw.line((x, y + 30, x - 35, y + 80), fill="#1C2833", width=4)
        draw.line((x, y + 30, x + 35, y + 80), fill="#1C2833", width=4)
        draw.text((x - 75, y + 90), label, font=font(24, True), fill="#111111")
    cases = [
        (620, 220, "Đăng ký và đăng nhập", 430),
        (1130, 220, "Tìm kiếm và xem sản phẩm", 480),
        (620, 400, "Quản lý giỏ hàng", 380),
        (1120, 400, "Đặt hàng và thanh toán VietQR", 520),
        (600, 650, "Quản lý danh mục sản phẩm", 480),
        (1130, 650, "Quản lý đơn hàng và kho", 480),
        (600, 830, "Quản lý tài khoản và phân quyền", 540),
        (1200, 830, "Báo cáo doanh thu", 390),
    ]
    for x, y, label, width in cases:
        draw.ellipse((x - width // 2, y - 55, x + width // 2, y + 55), fill="#EDF4FA", outline="#406E9F", width=3)
        tw = draw.textbbox((0, 0), label, font=font(21, True))[2]
        draw.text((x - tw // 2, y - 14), label, font=font(21, True), fill="#17365D")
    for end in [(405, 220), (890, 220), (430, 400), (860, 400)]:
        draw.line((205, 315, *end), fill="#7B8794", width=3)
    for end in [(360, 650), (890, 650), (330, 830), (1005, 830)]:
        draw.line((205, 755, *end), fill="#7B8794", width=3)
    path = IMAGES / "05-use-case.png"
    image.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[j])
            set_cell_shading(cell, NAVY if i == 0 else (PALE_BLUE if i % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(25, 25, 25)
                    run.bold = i == 0


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    style_table(table, widths)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_picture(doc, path, caption, width=6.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)
    c.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(25, 25, 25)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size in [("Title", 23), ("Heading 1", 16), ("Heading 2", 13)]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        p_pr = style._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def build_report():
    IMAGES.mkdir(parents=True, exist_ok=True)
    erd = build_erd()
    use_case = build_use_case()
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BÁO CÁO ĐỒ ÁN WEBSITE BÁN MỸ PHẨM SKINFOOD")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("PHP Laravel và MySQL")
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph("\n")
    add_table(doc, ["Thông tin", "Nội dung"], [
        ["Sinh viên", "........................................................"],
        ["Mã số sinh viên", "........................................................"],
        ["Lớp", "........................................................"],
        ["Giảng viên hướng dẫn", "........................................................"],
        ["Thời gian", "Năm học 2026"],
    ], [2.1, 4.5])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Mã nguồn chính: E:\\HK7\\PHP\\DACK\\laravel-app").italic = True
    doc.add_page_break()

    doc.add_heading("Tóm tắt kết quả", level=1)
    doc.add_paragraph(
        "Đồ án xây dựng website bán mỹ phẩm SkinFood bằng Laravel, Blade, Bootstrap và MySQL của Laragon. "
        "Hệ thống đã có luồng mua hàng, quản trị, nhập xuất kho, khuyến mãi, đánh giá và thanh toán VietQR. "
        "Dữ liệu mẫu gồm 25 sản phẩm, 7 danh mục, 12 loại sản phẩm, 9 thương hiệu, hai tài khoản mẫu và hai mã khuyến mãi."
    )
    add_table(doc, ["Nhóm yêu cầu", "Kết quả triển khai"], [
        ["Chức năng chính", "CRUD nhiều bảng có quan hệ, đăng nhập và phân quyền, tìm kiếm, phân trang, validation"],
        ["Laravel", "Eloquent ORM, migrations, seeder, controller, middleware và Blade Template"],
        ["Giao diện", "Bootstrap, bố cục desktop và responsive, trang người dùng và trang quản trị riêng"],
        ["Mở rộng", "Upload ảnh, AJAX giỏ hàng và mã giảm giá, VietQR API, webhook tự xác nhận thanh toán"],
        ["Kiểm thử", "PHP lint đạt, route Laravel hợp lệ và bộ test tự động đạt 3 trên 3"],
    ], [1.55, 5.05])

    doc.add_heading("Mục lục", level=1)
    add_bullets(doc, [
        "1 Tổng quan đề tài",
        "2 Phân tích chức năng và UML Use Case",
        "3 Thiết kế cơ sở dữ liệu",
        "4 Kiến trúc và triển khai Laravel",
        "5 Chức năng thanh toán VietQR",
        "6 Giao diện hệ thống",
        "7 Đối chiếu barem",
        "8 Hướng dẫn chạy và kiểm thử",
    ])
    doc.add_page_break()

    doc.add_heading("1 Tổng quan đề tài", level=1)
    doc.add_heading("1 1 Mục tiêu", level=2)
    doc.add_paragraph(
        "Website hỗ trợ khách hàng tìm kiếm mỹ phẩm, xem thông tin sản phẩm, quản lý giỏ hàng, áp mã giảm giá, "
        "đặt hàng và thanh toán. Quản trị viên quản lý danh mục, loại sản phẩm, thương hiệu, sản phẩm, nhà cung cấp, "
        "nhập xuất kho, đơn hàng, tài khoản, đánh giá, liên hệ và báo cáo doanh thu."
    )
    doc.add_heading("1 2 Công nghệ", level=2)
    add_table(doc, ["Thành phần", "Công nghệ", "Vai trò"], [
        ["Backend", "PHP 8.3 và Laravel", "Routing, controller, middleware, validation và nghiệp vụ"],
        ["Dữ liệu", "MySQL Laragon", "Lưu dữ liệu quan hệ, khóa ngoại và giao dịch"],
        ["Frontend", "Blade, Bootstrap và JavaScript", "Hiển thị responsive và tương tác AJAX"],
        ["Thanh toán", "VietQR Quick Link và webhook", "Tạo QR đúng số tiền và nhận trạng thái giao dịch"],
    ], [1.2, 2.0, 3.4])

    doc.add_heading("2 Phân tích chức năng và UML Use Case", level=1)
    add_picture(doc, use_case, "Hình 1 UML Use Case của hệ thống", 6.8)
    add_table(doc, ["Tác nhân", "Chức năng chính"], [
        ["Khách hàng", "Đăng ký, đăng nhập, tìm kiếm, xem sản phẩm, giỏ hàng, mã giảm giá, đặt hàng, thanh toán và đánh giá"],
        ["Quản trị viên", "CRUD dữ liệu, quản lý kho và đơn hàng, duyệt đánh giá, quản lý tài khoản và xem báo cáo"],
        ["Dịch vụ thanh toán", "Tạo ảnh VietQR và gửi webhook giao dịch thành công về hệ thống"],
    ], [1.45, 5.15])

    doc.add_heading("3 Thiết kế cơ sở dữ liệu", level=1)
    doc.add_paragraph(
        "Cơ sở dữ liệu dùng khóa chính số nguyên, khóa ngoại và các bảng chi tiết có khóa chính ghép. "
        "Các quan hệ chính bảo đảm danh mục và loại sản phẩm liên kết với sản phẩm; người dùng liên kết với tài khoản và đơn hàng; "
        "đơn hàng liên kết với chi tiết đơn hàng và khuyến mãi."
    )
    add_picture(doc, erd, "Hình 2 ERD rút gọn của cơ sở dữ liệu SkinFood", 6.8)
    add_table(doc, ["Nhóm bảng", "Bảng tiêu biểu", "Quan hệ"], [
        ["Danh mục sản phẩm", "DanhMuc, LoaiSP, ThuongHieu, SanPham", "Một nhiều qua MaDM, MaLoai và MaTH"],
        ["Người dùng", "NguoiDung, TaiKhoan", "Một một qua MaND"],
        ["Bán hàng", "DonHang, ChiTietDonHangs, KhuyenMai", "Một nhiều và nhiều nhiều"],
        ["Kho", "PhieuNhap, ChiTietPhieuNhap, PhieuXuat, ChiTietPhieuXuat", "Phiếu và dòng chi tiết"],
    ], [1.35, 2.6, 2.65])

    doc.add_heading("4 Kiến trúc và triển khai Laravel", level=1)
    doc.add_heading("4 1 Tổ chức mã nguồn", level=2)
    doc.add_paragraph(
        "Controller xử lý yêu cầu và nghiệp vụ; Model ánh xạ bảng MySQL bằng Eloquent; Blade render giao diện; middleware auth và admin bảo vệ khu vực quản trị. "
        "VietQrService tách riêng việc tạo payload chuẩn VietQR và Quick Link của API bên ngoài."
    )
    doc.add_heading("4 2 CRUD validation và dữ liệu quan hệ", level=2)
    add_bullets(doc, [
        "CRUD đầy đủ cho danh mục, loại sản phẩm, thương hiệu, nhà cung cấp, tài khoản, khuyến mãi và sản phẩm.",
        "Form dùng Laravel validation cho chuỗi, số, ngày, email, khóa ngoại, ảnh và quy tắc nghiệp vụ.",
        "Phiếu nhập và phiếu xuất dùng transaction và lockForUpdate để cập nhật tồn kho an toàn.",
        "Các truy vấn danh sách dùng paginate và eager loading để giảm dữ liệu và số lần truy vấn.",
    ])
    doc.add_heading("4 3 Đăng nhập và phân quyền", level=2)
    doc.add_paragraph(
        "Hệ thống có hai vai trò Admin và User. Mật khẩu tài khoản mới được băm, tài khoản có trạng thái hoạt động hoặc khóa. "
        "AdminMiddleware từ chối truy cập khu vực quản trị khi người dùng không có vai trò Admin."
    )

    doc.add_heading("5 Chức năng thanh toán VietQR", level=1)
    doc.add_paragraph(
        "Khi khách chọn chuyển khoản, hệ thống tạo mã tham chiếu mới và đặt thời hạn 15 phút. QR chứa đúng ngân hàng, số tài khoản, "
        "tổng tiền và nội dung đơn hàng. Khi hết thời hạn, trình duyệt xóa QR và khóa nút xác nhận; khi người dùng mở lại trang, hệ thống tạo QR mới."
    )
    add_bullets(doc, [
        "VietQR Quick Link tạo ảnh QR qua dịch vụ bên ngoài; payload EMV cục bộ là phương án dự phòng khi mất mạng.",
        "Trang thanh toán kiểm tra trạng thái ba giây một lần và chỉ ẩn QR khi đơn đã được xác nhận thanh toán.",
        "Webhook yêu cầu secret, đối chiếu mã tham chiếu, thời hạn và số tiền trước khi cập nhật trạng thái Đã thanh toán.",
        "Nút Tôi đã chuyển khoản chỉ chuyển đơn sang trạng thái chờ xác nhận, không tự đánh dấu đã thanh toán.",
    ])
    add_picture(doc, IMAGES / "03-thanh-toan-vietqr.png", "Hình 3 Thanh toán chuyển khoản với QR động và bộ đếm 15 phút", 6.8)

    doc.add_heading("6 Giao diện hệ thống", level=1)
    add_picture(doc, IMAGES / "01-trang-chu.png", "Hình 4 Trang chủ và danh sách sản phẩm", 6.8)
    add_picture(doc, IMAGES / "02-quan-tri.png", "Hình 5 Dashboard dành cho quản trị viên", 6.8)
    doc.add_paragraph(
        "Giao diện dùng Bootstrap, lưới sản phẩm và thanh điều hướng có quy tắc responsive. Trang quản trị tách menu theo nhóm kho hàng, kinh doanh và tài khoản."
    )

    doc.add_page_break()
    doc.add_heading("7 Đối chiếu barem", level=1)
    rows = [
        ["CRUD ít nhất hai bảng có quan hệ", "Đạt", "Danh mục, loại, sản phẩm, thương hiệu, nhà cung cấp, khuyến mãi"],
        ["Đăng nhập đăng xuất và hai vai trò", "Đạt", "Admin, User, middleware và trạng thái tài khoản"],
        ["Tìm kiếm và phân trang", "Đạt", "Tìm sản phẩm và paginate ở trang khách và quản trị"],
        ["Validation form", "Đạt", "Laravel validate cho các form chính"],
        ["Eloquent ORM", "Đạt", "Model và relationship cho dữ liệu nghiệp vụ"],
        ["Migration và Seeder", "Đạt", "Schema đầy đủ và SkinFoodSeeder"],
        ["Controller Middleware Blade", "Đạt", "Tổ chức đúng kiến trúc Laravel"],
        ["Bootstrap và responsive", "Đạt", "Trang khách và trang quản trị"],
        ["Báo cáo ERD UML ảnh", "Đạt", "Tài liệu này có ERD, UML và ảnh chụp thật"],
        ["Upload file", "Đạt", "Upload và thay ảnh sản phẩm"],
        ["AJAX", "Đạt", "Giỏ hàng, mã giảm giá và polling thanh toán"],
        ["API bên ngoài", "Đạt", "VietQR Quick Link"],
        ["Tối ưu", "Đạt", "Transaction, lock, eager loading và pagination"],
    ]
    add_table(doc, ["Tiêu chí", "Kết quả", "Minh chứng"], rows, [2.45, 0.75, 3.4])
    doc.add_paragraph(
        "Phần mã nguồn và báo cáo đã đáp ứng các tiêu chí có thể kiểm tra trong project. Điểm thuyết trình và vấn đáp phụ thuộc phần trình bày trực tiếp của sinh viên."
    )

    doc.add_heading("8 Hướng dẫn chạy và kiểm thử", level=1)
    add_table(doc, ["Bước", "Lệnh hoặc thao tác"], [
        ["1", "Mở Laragon và khởi động MySQL"],
        ["2", "cd E:\\HK7\\PHP\\DACK\\laravel-app"],
        ["3", "php artisan migrate --seed"],
        ["4", "php artisan serve --port=8001"],
        ["5", "Mở http://127.0.0.1:8001"],
        ["6", "Chạy kiểm thử bằng php artisan test"],
    ], [0.7, 5.9])
    doc.add_heading("8 1 Tài khoản mẫu", level=2)
    add_table(doc, ["Vai trò", "Tên đăng nhập", "Mật khẩu"], [
        ["Admin", "admin@skinfood.com", "123456"],
        ["User", "user@skinfood.com", "123456"],
    ], [1.25, 3.5, 1.85])
    doc.add_heading("8 2 Cấu hình thanh toán", level=2)
    doc.add_paragraph(
        "Thông tin BANK_BIN, BANK_ACCOUNT_NUMBER, BANK_ACCOUNT_NAME và BANK_WEBHOOK_SECRET được đặt trong file .env. "
        "Khi triển khai thật, cần đổi webhook secret và cấu hình nhà cung cấp dữ liệu giao dịch ngân hàng gọi endpoint POST /api/thanh-toan/webhook."
    )
    doc.add_heading("Tài liệu tham khảo", level=1)
    add_bullets(doc, [
        "Laravel Documentation https://laravel.com/docs",
        "VietQR Quick Link https://www.vietqr.io/qr-dong",
        "Bootstrap Documentation https://getbootstrap.com/docs",
    ])

    doc.core_properties.title = "Báo cáo đồ án website bán mỹ phẩm SkinFood"
    doc.core_properties.subject = "PHP Laravel MySQL"
    doc.core_properties.author = "Nhóm đồ án SkinFood"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
